import os
from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from careerradar.profile.models import (
    ResumeEducation,
    ResumeRole,
    ResumeSkillCategory,
    ResumeSubsection,
    TailoredResumePayload,
)
from careerradar.profile.renderer import convert_to_pdf, render_docx, verify_page_count


def test_render_docx_and_convert_pdf(tmp_path: Path):
    payload = TailoredResumePayload(
        name="Jane Doe",
        contact_line_1="San Francisco Bay Area | jane.doe@example.com | 555-019-2834",
        contact_line_2="github.com/janedoe | linkedin.com/in/janedoe",
        summary=(
            "Senior full-stack software engineer specialized in distributed systems, "
            "LLM orchestration, and modern cloud infrastructure."
        ),
        experience=[
            ResumeRole(
                title="Lead Software Engineer",
                company="Acme Cloud Systems",
                dates="Jan 2020 - present",
                subsections=[
                    ResumeSubsection(
                        heading="Designed platform infrastructure and web services:",
                        bullets=[
                            "Architected distributed web application handling 50k+ users.",
                            "Automated cloud migrations saving 30+ engineering hrs/week.",
                        ],
                    ),
                    ResumeSubsection(
                        heading="Data Pipeline Modernization:",
                        bullets=[
                            "Built streaming pipeline processing 10M events daily.",
                        ],
                    ),
                ],
            ),
            ResumeRole(
                title="Senior Backend Engineer",
                company="Beta Corp",
                dates="Jan 2018 - Dec 2019",
                subsections=[
                    ResumeSubsection(
                        heading="Core APIs:",
                        bullets=[
                            "Developed high-throughput REST microservices in Python.",
                        ],
                    )
                ],
            ),
        ],
        skills=[
            ResumeSkillCategory(
                category="Infrastructure & Cloud",
                skills="AWS, Docker, Terraform, Kubernetes, Linux",
            ),
            ResumeSkillCategory(
                category="Full-Stack & Languages",
                skills="Python, TypeScript, FastAPI, React, Node.js, SQLite, PostgreSQL",
            ),
        ],
        education=[
            ResumeEducation(
                institution="State University",
                degree="BS in Computer Science",
            )
        ],
    )

    docx_path = str(tmp_path / "test_resume.docx")
    out_file = render_docx(payload, docx_path)
    assert os.path.exists(out_file)
    assert os.path.getsize(out_file) > 1000

    # Inspect docx contents
    doc = docx.Document(docx_path)
    text = " ".join(p.text for p in doc.paragraphs)
    assert "JANE DOE" in text or "Jane Doe" in text
    assert "Acme Cloud Systems" in text
    assert "EXPERIENCE" in text
    assert "SKILLS" in text
    assert "EDUCATION" in text

    # Verify section margins: 0.75 in top, bottom, left, right
    for section in doc.sections:
        assert section.top_margin == Inches(0.75)
        assert section.bottom_margin == Inches(0.75)
        assert section.left_margin == Inches(0.75)
        assert section.right_margin == Inches(0.75)

    # Verify contact header and main headers have page-wide bottom borders
    headers_with_border = []
    for p in doc.paragraphs:
        if p.text in [payload.contact_line_2, "EXPERIENCE", "SKILLS", "EDUCATION"]:
            pPr = p._element.pPr
            pbdr = pPr.find(qn("w:pBdr")) if pPr is not None else None
            if pbdr is not None and pbdr.find(qn("w:bottom")) is not None:
                headers_with_border.append(p.text)
    assert headers_with_border == [payload.contact_line_2, "EXPERIENCE", "SKILLS", "EDUCATION"]

    # Verify summary spacing after contact header
    sum_p = [p for p in doc.paragraphs if "Senior full-stack software engineer" in p.text]
    assert len(sum_p) == 1
    assert sum_p[0].paragraph_format.space_before == Pt(6)

    # Verify subsection and role spacing (1 line = 10pt space_before on subsequent items)
    sub2_p = [p for p in doc.paragraphs if "Data Pipeline Modernization" in p.text]
    assert len(sub2_p) == 1
    assert sub2_p[0].paragraph_format.space_before == Pt(10)

    role2_p = [p for p in doc.paragraphs if "Beta Corp" in p.text]
    assert len(role2_p) == 1
    assert role2_p[0].paragraph_format.space_before == Pt(10)

    # Test PDF conversion
    pdf_path = convert_to_pdf(docx_path, str(tmp_path))
    if pdf_path:
        assert os.path.exists(pdf_path)
        pages = verify_page_count(pdf_path)
        assert pages == 1
