"""DOCX and PDF rendering engine for 1-page tailored resumes."""

import os
import subprocess
from pathlib import Path
from typing import Any

import docx
import pypdf
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from pypdf.errors import PdfReadError

from careerradar.core.logger import get_logger
from careerradar.core.paths import GENERATED_RESUMES_DIR, TEMPLATES_DIR
from careerradar.resumes.models import TailoredResumePayload

logger = get_logger()

DEFAULT_TEMPLATE_PATH = os.path.join(TEMPLATES_DIR, "resume_template.docx")
FALLBACK_TEMPLATE_PATH = "/opt/Downloads/resume.docx"


def _get_template_path() -> str:
    if os.path.exists(DEFAULT_TEMPLATE_PATH):
        return DEFAULT_TEMPLATE_PATH
    if os.path.exists(FALLBACK_TEMPLATE_PATH):
        return FALLBACK_TEMPLATE_PATH
    raise FileNotFoundError("No resume docx template found in templates/ or downloads.")


def _add_right_tab(paragraph: Any) -> None:
    pPr = paragraph._element.get_or_add_pPr()
    xml_str = (
        '<w:tabs xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:tab w:val="right" w:pos="9360"/>'
        "</w:tabs>"
    )
    tabs = parse_xml(xml_str)
    pPr.append(tabs)


def render_docx(
    payload: TailoredResumePayload,
    output_path: str,
    template_path: str | None = None,
) -> str:
    """Render the tailored resume into a high-fidelity 1-page DOCX file."""
    tpl = template_path or _get_template_path()
    doc = docx.Document(tpl)

    # Ensure section margins: 0.5 in top/bottom, 1.0 in left/right
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Clear body paragraphs while preserving styles, fontTable, settings
    body = doc._body._element
    for p in body.findall(qn("w:p")):
        body.remove(p)

    # 1. Header
    p_name = doc.add_paragraph(style="Heading 1")
    p_name.paragraph_format.space_before = Pt(0)
    p_name.paragraph_format.space_after = Pt(0)
    p_name.paragraph_format.line_spacing = 1.0
    r_name = p_name.add_run(payload.name)
    r_name.font.size = Pt(24)
    r_name.font.name = "Roboto"

    p_c1 = doc.add_paragraph()
    p_c1.paragraph_format.space_before = Pt(0)
    p_c1.paragraph_format.space_after = Pt(0)
    p_c1.paragraph_format.line_spacing = 1.0
    r_c1 = p_c1.add_run(payload.contact_line_1)
    r_c1.font.size = Pt(10.5)
    r_c1.font.name = "Roboto"

    p_c2 = doc.add_paragraph()
    p_c2.paragraph_format.space_before = Pt(0)
    p_c2.paragraph_format.space_after = Pt(4)
    p_c2.paragraph_format.line_spacing = 1.0
    r_c2 = p_c2.add_run(payload.contact_line_2)
    r_c2.font.size = Pt(10.5)
    r_c2.font.name = "Roboto"

    # 2. Summary
    p_sum = doc.add_paragraph(style="Subtitle")
    p_sum.paragraph_format.space_before = Pt(2)
    p_sum.paragraph_format.space_after = Pt(4)
    p_sum.paragraph_format.line_spacing = 1.15
    r_sum = p_sum.add_run(payload.summary)
    r_sum.font.size = Pt(10.5)
    r_sum.font.name = "Roboto"

    # 3. EXPERIENCE Section
    p_exp_h = doc.add_paragraph(style="Heading 2")
    p_exp_h.paragraph_format.space_before = Pt(6)
    p_exp_h.paragraph_format.space_after = Pt(2)
    p_exp_h.paragraph_format.line_spacing = 1.0
    r_exp_h = p_exp_h.add_run("EXPERIENCE")
    r_exp_h.bold = True
    r_exp_h.font.size = Pt(11)
    r_exp_h.font.name = "Roboto"

    for role in payload.experience:
        p_role = doc.add_paragraph(style="Heading 3")
        p_role.paragraph_format.space_before = Pt(4)
        p_role.paragraph_format.space_after = Pt(1)
        _add_right_tab(p_role)
        r_role = p_role.add_run(f"{role.title}, {role.company}\t{role.dates}")
        r_role.font.name = "Roboto Medium"
        r_role.font.size = Pt(10.5)

        for sub in role.subsections:
            if sub.heading and sub.heading.strip():
                p_sub = doc.add_paragraph()
                p_sub.paragraph_format.space_before = Pt(2)
                p_sub.paragraph_format.space_after = Pt(1)
                p_sub.paragraph_format.line_spacing = 1.05
                r_sub = p_sub.add_run(sub.heading)
                r_sub.italic = True
                r_sub.font.size = Pt(10.5)
                r_sub.font.name = "Roboto"

            for bullet in sub.bullets:
                bp = doc.add_paragraph()
                bp.paragraph_format.left_indent = Inches(0.25)
                bp.paragraph_format.first_line_indent = Inches(-0.15)
                bp.paragraph_format.space_before = Pt(0)
                bp.paragraph_format.space_after = Pt(1.5)
                bp.paragraph_format.line_spacing = 1.05
                br = bp.add_run("•  " + bullet)
                br.font.size = Pt(10)
                br.font.name = "Roboto"

    # 4. SKILLS Section
    p_sk_h = doc.add_paragraph(style="Heading 2")
    p_sk_h.paragraph_format.space_before = Pt(5)
    p_sk_h.paragraph_format.space_after = Pt(2)
    p_sk_h.paragraph_format.line_spacing = 1.0
    r_sk_h = p_sk_h.add_run("SKILLS")
    r_sk_h.bold = True
    r_sk_h.font.size = Pt(11)
    r_sk_h.font.name = "Roboto"

    for cat in payload.skills:
        p_sk = doc.add_paragraph()
        p_sk.paragraph_format.space_before = Pt(0)
        p_sk.paragraph_format.space_after = Pt(1.5)
        p_sk.paragraph_format.line_spacing = 1.05
        r_cat = p_sk.add_run(f"{cat.category}: ")
        r_cat.bold = True
        r_cat.font.size = Pt(10)
        r_cat.font.name = "Roboto"
        r_items = p_sk.add_run(cat.skills)
        r_items.font.size = Pt(10)
        r_items.font.name = "Roboto"

    # 5. EDUCATION Section
    p_ed_h = doc.add_paragraph(style="Heading 2")
    p_ed_h.paragraph_format.space_before = Pt(5)
    p_ed_h.paragraph_format.space_after = Pt(2)
    p_ed_h.paragraph_format.line_spacing = 1.0
    r_ed_h = p_ed_h.add_run("EDUCATION")
    r_ed_h.bold = True
    r_ed_h.font.size = Pt(11)
    r_ed_h.font.name = "Roboto"

    for edu in payload.education:
        p_ed = doc.add_paragraph()
        p_ed.paragraph_format.space_before = Pt(0)
        p_ed.paragraph_format.space_after = Pt(1.5)
        p_ed.paragraph_format.line_spacing = 1.05
        r_ed = p_ed.add_run(f"{edu.institution}, {edu.degree}")
        r_ed.bold = True
        r_ed.font.size = Pt(10)
        r_ed.font.name = "Roboto"

    out_p = Path(output_path)
    out_p.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_p))
    logger.info("Rendered resume docx saved to %s", output_path)
    return str(out_p)


def convert_to_pdf(docx_path: str, output_dir: str | None = None) -> str | None:
    """Convert a docx file to PDF using headless LibreOffice."""
    out_dir = output_dir or os.path.dirname(docx_path) or GENERATED_RESUMES_DIR
    Path(out_dir).mkdir(parents=True, exist_ok=True)
    try:
        cmd = [
            "libreoffice",
            "--headless",
            "--convert-to",
            "pdf",
            docx_path,
            "--outdir",
            out_dir,
        ]
        subprocess.run(cmd, check=True, capture_output=True)
        stem = Path(docx_path).stem
        pdf_path = os.path.join(out_dir, f"{stem}.pdf")
        if os.path.exists(pdf_path):
            return pdf_path
        logger.warning("PDF conversion succeeded but file not found at %s", pdf_path)
        return None
    except (subprocess.SubprocessError, OSError) as exc:
        logger.warning("LibreOffice PDF conversion failed for %s: %s", docx_path, exc)
        return None


def verify_page_count(pdf_path: str) -> int:
    """Check page count of rendered PDF."""
    try:
        reader = pypdf.PdfReader(pdf_path)
        return len(reader.pages)
    except (PdfReadError, OSError) as exc:
        logger.warning("Failed reading PDF page count for %s: %s", pdf_path, exc)
        return 0
