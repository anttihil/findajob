"""DOCX and PDF rendering engine for 1-page tailored resumes."""

import os
import subprocess
from pathlib import Path
from typing import Any

import docx
import pypdf
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from pypdf.errors import PdfReadError

from careerradar.core.logger import get_logger
from careerradar.core.paths import GENERATED_RESUMES_DIR, TEMPLATES_DIR
from careerradar.profile.models import TailoredResumePayload

logger = get_logger()

DEFAULT_TEMPLATE_PATH = os.path.join(TEMPLATES_DIR, "resume_template.docx")


def _get_template_path() -> str | None:
    if os.path.exists(DEFAULT_TEMPLATE_PATH):
        return DEFAULT_TEMPLATE_PATH
    fallback = os.path.expanduser("~/Downloads/resume.docx")
    if os.path.exists(fallback):
        return fallback
    return None


def _add_right_tab(paragraph: Any) -> None:
    pPr = paragraph._element.get_or_add_pPr()
    xml_str = (
        '<w:tabs xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:tab w:val="right" w:pos="10080"/>'
        "</w:tabs>"
    )
    tabs = parse_xml(xml_str)
    pPr.append(tabs)


def _add_bottom_border(
    paragraph: Any,
    sz: str = "6",
    color: str = "auto",
    space: str = "1",
) -> None:
    pPr = paragraph._element.get_or_add_pPr()
    existing_pBdr = pPr.find(qn("w:pBdr"))
    if existing_pBdr is not None:
        pPr.remove(existing_pBdr)
    xml_str = (
        f'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f'<w:bottom w:val="single" w:sz="{sz}" w:space="{space}" w:color="{color}"/>'
        f"</w:pBdr>"
    )
    pBdr = parse_xml(xml_str)
    pPr.append(pBdr)


def render_docx(
    payload: TailoredResumePayload,
    output_path: str,
    template_path: str | None = None,
) -> str:
    """Render the tailored resume into a high-fidelity 1-page DOCX file."""
    tpl = template_path or _get_template_path()
    doc = docx.Document(tpl) if (tpl and os.path.exists(tpl)) else docx.Document()

    # Ensure section margins: 0.75 in all around
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Clear body paragraphs while preserving styles, fontTable, settings
    body = doc._body._element
    for p in body.findall(qn("w:p")):
        body.remove(p)

    # 1. Header
    p_name = doc.add_paragraph(style="Heading 1")
    p_name.paragraph_format.space_before = Pt(0)
    p_name.paragraph_format.space_after = Pt(0)
    p_name.paragraph_format.line_spacing = 1.0
    r_name = p_name.add_run(payload.name)
    r_name.font.size = Pt(24)
    r_name.font.name = "Arial"

    p_c1 = doc.add_paragraph()
    p_c1.paragraph_format.space_before = Pt(0)
    p_c1.paragraph_format.space_after = Pt(0)
    p_c1.paragraph_format.line_spacing = 1.0
    r_c1 = p_c1.add_run(payload.contact_line_1)
    r_c1.font.size = Pt(10.5)
    r_c1.font.name = "Arial"

    last_contact_p = p_c1
    if payload.contact_line_2 and payload.contact_line_2.strip():
        p_c2 = doc.add_paragraph()
        p_c2.paragraph_format.space_before = Pt(0)
        p_c2.paragraph_format.space_after = Pt(4)
        p_c2.paragraph_format.line_spacing = 1.0
        r_c2 = p_c2.add_run(payload.contact_line_2)
        r_c2.font.size = Pt(10.5)
        r_c2.font.name = "Arial"
        last_contact_p = p_c2

    _add_bottom_border(last_contact_p)

    # 2. Summary
    p_sum = doc.add_paragraph(style="Subtitle")
    p_sum.paragraph_format.space_before = Pt(6)
    p_sum.paragraph_format.space_after = Pt(4)
    p_sum.paragraph_format.line_spacing = 1.15
    r_sum = p_sum.add_run(payload.summary)
    r_sum.font.size = Pt(10.5)
    r_sum.font.name = "Arial"

    # 3. EXPERIENCE Section
    p_exp_h = doc.add_paragraph(style="Heading 2")
    p_exp_h.paragraph_format.space_before = Pt(10)
    p_exp_h.paragraph_format.space_after = Pt(2)
    p_exp_h.paragraph_format.line_spacing = 1.0
    _add_bottom_border(p_exp_h)
    r_exp_h = p_exp_h.add_run("EXPERIENCE")
    r_exp_h.bold = True
    r_exp_h.font.size = Pt(11)
    r_exp_h.font.name = "Arial"

    for role_idx, role in enumerate(payload.experience):
        p_role = doc.add_paragraph(style="Heading 3")
        p_role.paragraph_format.space_before = Pt(10) if role_idx > 0 else Pt(4)
        p_role.paragraph_format.space_after = Pt(1)
        _add_right_tab(p_role)
        r_role = p_role.add_run(f"{role.title}, {role.company}\t{role.dates}")
        r_role.font.name = "Arial"
        r_role.font.size = Pt(10.5)

        for sub_idx, sub in enumerate(role.subsections):
            if sub.heading and sub.heading.strip():
                p_sub = doc.add_paragraph()
                p_sub.paragraph_format.space_before = Pt(10) if sub_idx > 0 else Pt(2)
                p_sub.paragraph_format.space_after = Pt(1)
                p_sub.paragraph_format.line_spacing = 1.05
                r_sub = p_sub.add_run(sub.heading)
                r_sub.italic = True
                r_sub.font.size = Pt(10.5)
                r_sub.font.name = "Arial"

            for b_idx, bullet in enumerate(sub.bullets):
                bp = doc.add_paragraph()
                bp.paragraph_format.left_indent = Inches(0.25)
                bp.paragraph_format.first_line_indent = Inches(-0.15)
                if not (sub.heading and sub.heading.strip()) and sub_idx > 0 and b_idx == 0:
                    bp.paragraph_format.space_before = Pt(10)
                else:
                    bp.paragraph_format.space_before = Pt(0)
                bp.paragraph_format.space_after = Pt(1.5)
                bp.paragraph_format.line_spacing = 1.05
                br = bp.add_run("•\t" + bullet)
                br.font.size = Pt(10)
                br.font.name = "Arial"

    # 4. SKILLS Section
    p_sk_h = doc.add_paragraph(style="Heading 2")
    p_sk_h.paragraph_format.space_before = Pt(10)
    p_sk_h.paragraph_format.space_after = Pt(2)
    p_sk_h.paragraph_format.line_spacing = 1.0
    _add_bottom_border(p_sk_h)
    r_sk_h = p_sk_h.add_run("SKILLS")
    r_sk_h.bold = True
    r_sk_h.font.size = Pt(11)
    r_sk_h.font.name = "Arial"

    for cat in payload.skills:
        p_sk = doc.add_paragraph()
        p_sk.paragraph_format.space_before = Pt(0)
        p_sk.paragraph_format.space_after = Pt(1.5)
        p_sk.paragraph_format.line_spacing = 1.05
        r_cat = p_sk.add_run(f"{cat.category}: ")
        r_cat.bold = True
        r_cat.font.size = Pt(10)
        r_cat.font.name = "Arial"
        r_items = p_sk.add_run(cat.skills)
        r_items.font.size = Pt(10)
        r_items.font.name = "Arial"

    # 5. EDUCATION Section
    p_ed_h = doc.add_paragraph(style="Heading 2")
    p_ed_h.paragraph_format.space_before = Pt(10)
    p_ed_h.paragraph_format.space_after = Pt(2)
    p_ed_h.paragraph_format.line_spacing = 1.0
    _add_bottom_border(p_ed_h)
    r_ed_h = p_ed_h.add_run("EDUCATION")
    r_ed_h.bold = True
    r_ed_h.font.size = Pt(11)
    r_ed_h.font.name = "Arial"

    for edu in payload.education:
        p_ed = doc.add_paragraph()
        p_ed.paragraph_format.space_before = Pt(0)
        p_ed.paragraph_format.space_after = Pt(1.5)
        p_ed.paragraph_format.line_spacing = 1.05
        r_ed = p_ed.add_run(f"{edu.institution}, {edu.degree}")
        r_ed.bold = True
        r_ed.font.size = Pt(10)
        r_ed.font.name = "Arial"

    out_p = Path(output_path)
    out_p.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_p))
    logger.info("Rendered resume docx saved to %s", output_path)
    return str(out_p)


def escape_typst(text: str) -> str:
    """Escape special Typst markup characters in user-provided text."""
    if not text:
        return ""
    replacements = [
        ("\\", "\\\\"),
        ("@", "\\@"),
        ("<", "\\<"),
        (">", "\\>"),
        ("$", "\\$"),
        ("#", "\\#"),
        ("[", "\\["),
        ("]", "\\]"),
    ]
    for old, new in replacements:
        text = text.replace(old, new)
    return text


def generate_typst_source(payload: TailoredResumePayload) -> str:
    """Generate high-fidelity 1-page Typst markup source code from resume payload."""
    lines: list[str] = [
        "// ============================================================================",
        f"// Tailored 1-Page Resume: {escape_typst(payload.name)}",
        "// Generated with Typst - Fast, deterministic typesetting engine",
        "// ============================================================================",
        "",
        "#set page(",
        '  paper: "us-letter",',
        "  margin: (x: 0.5in, top: 0.5in, bottom: 0.5in),",
        ")",
        "#set text(",
        '  font: ("Arial", "Liberation Sans", "DejaVu Sans"),',
        "  size: 11pt,",
        '  fill: rgb("#000000"),',
        ")",
        "#set par(justify: false, leading: 0.50em)",
        "",
        "// Helper: Section heading with bottom rule (tight underline)",
        "#let section-heading(title) = {",
        "  block(",
        "    width: 100%,",
        '    stroke: (bottom: 1pt + rgb("#000000")),',
        "    inset: (bottom: 2pt),",
        "    above: 10pt,",
        "    below: 10pt,",
        '  )[#text(11pt, weight: "bold")[#upper(title)]]',
        "}",
        "",
        "// --- Header ---",
        f'#text(20pt, weight: "bold")[{escape_typst(payload.name)}]',
        "#v(1pt)",
        f"#text()[{escape_typst(payload.contact_line_1)}]",
    ]
    if payload.contact_line_2 and payload.contact_line_2.strip():
        lines.append(f"#v(1pt)\n#text()[{escape_typst(payload.contact_line_2)}]")
    lines.extend(
        [
            "",
            "#v(-3pt)",
            '#line(length: 100%, stroke: 1pt + rgb("#000000"))',
            "#v(1pt)",
            "",
            "// --- Summary ---",
            f"#text()[{escape_typst(payload.summary)}]",
            "",
            "// --- Experience ---",
            '#section-heading("EXPERIENCE")',
        ]
    )

    for r_idx, role in enumerate(payload.experience):
        if r_idx > 0:
            lines.append("#v(4pt)")
        lines.extend(
            [
                "#grid(",
                "  columns: (1fr, auto),",
                f"  [*{escape_typst(role.title)}*, {escape_typst(role.company)}],",
                f'  text(9pt, fill: rgb("#000000"))[{escape_typst(role.dates)}],',
                ")",
            ]
        )
        for sub in role.subsections:
            if sub.heading and sub.heading.strip():
                lines.extend(
                    [
                        "#v(1pt)",
                        f'#text(weight: "bold")[{escape_typst(sub.heading)}]',
                    ]
                )
            lines.append("#v(1pt)")
            lines.append("#list(")
            lines.append("  tight: true,")
            for bullet in sub.bullets:
                lines.append(f"  [{escape_typst(bullet)}],")
            lines.append(")")

    lines.extend(
        [
            "",
            "// --- Skills ---",
            '#section-heading("SKILLS")',
        ]
    )
    for cat in payload.skills:
        lines.append(f"*{escape_typst(cat.category)}:* {escape_typst(cat.skills)} \\")

    lines.extend(
        [
            "",
            "// --- Education ---",
            '#section-heading("EDUCATION")',
        ]
    )
    for edu in payload.education:
        lines.extend(
            [
                "#grid(",
                "  columns: (1fr, auto),",
                f"  [*{escape_typst(edu.institution)}*, {escape_typst(edu.degree)}],",
                "  [],",
                ")",
            ]
        )

    return "\n".join(lines)


def render_typst(
    payload: TailoredResumePayload,
    output_path: str,
) -> str:
    """Render the tailored resume into a high-fidelity 1-page Typst (.typ) markup file."""
    source = generate_typst_source(payload)
    out_p = Path(output_path)
    out_p.parent.mkdir(parents=True, exist_ok=True)
    out_p.write_text(source, encoding="utf-8")
    logger.info("Rendered resume typst saved to %s", output_path)
    return str(out_p)


def compile_typst_to_pdf(
    typst_path: str,
    output_dir: str | None = None,
    output_pdf_path: str | None = None,
) -> str | None:
    """Compile a .typ file to PDF using the native typst compiler."""
    import typst

    if output_pdf_path:
        pdf_path = output_pdf_path
    else:
        out_dir = output_dir or os.path.dirname(typst_path) or GENERATED_RESUMES_DIR
        stem = Path(typst_path).stem
        pdf_path = os.path.join(out_dir, f"{stem}.pdf")

    Path(pdf_path).parent.mkdir(parents=True, exist_ok=True)
    try:
        typst.compile(typst_path, output=pdf_path)
        if os.path.exists(pdf_path):
            logger.info("Typst compiled %s -> %s", typst_path, pdf_path)
            return pdf_path
        logger.warning("Typst compilation succeeded but file not found at %s", pdf_path)
        return None
    except (OSError, RuntimeError) as exc:
        logger.warning("Typst PDF compilation failed for %s: %s", typst_path, exc)
        return None


def convert_to_pdf(source_path: str, output_dir: str | None = None) -> str | None:
    """Convert a resume source file (.typ or .docx) to PDF."""
    if source_path.endswith(".typ"):
        return compile_typst_to_pdf(source_path, output_dir=output_dir)

    # Fallback to headless LibreOffice for legacy .docx files
    docx_path = source_path
    out_dir = output_dir or os.path.dirname(docx_path) or GENERATED_RESUMES_DIR
    Path(out_dir).mkdir(parents=True, exist_ok=True)
    try:
        cmd = [
            "libreoffice",
            "--headless",
            "--convert-to",
            "pdf",
            docx_path,
            "--outdir",
            out_dir,
        ]
        subprocess.run(cmd, check=True, capture_output=True)
        stem = Path(docx_path).stem
        pdf_path = os.path.join(out_dir, f"{stem}.pdf")
        if os.path.exists(pdf_path):
            return pdf_path
        logger.warning("PDF conversion succeeded but file not found at %s", pdf_path)
        return None
    except (subprocess.SubprocessError, OSError) as exc:
        logger.warning("LibreOffice PDF conversion failed for %s: %s", docx_path, exc)
        return None


def verify_page_count(pdf_path: str) -> int:
    """Check page count of rendered PDF."""
    try:
        reader = pypdf.PdfReader(pdf_path)
        return len(reader.pages)
    except (PdfReadError, OSError) as exc:
        logger.warning("Failed reading PDF page count for %s: %s", pdf_path, exc)
        return 0
