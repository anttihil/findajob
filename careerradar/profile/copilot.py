"""AI Copilot and resume parsing engine for the unified Profile.

Supports:
1. Parsing uploaded resume files (.pdf, .docx, .txt, .md).
2. LLM-based extraction into structured Profile.
3. Interactive Copilot chat to refine, question, rewrite, and update the Profile.
"""

import io

from pydantic import BaseModel, Field

from careerradar.core.llm import DEFAULT_AGENT_MODEL, invoke_structured, structured_model
from careerradar.core.logger import get_logger
from careerradar.profile.models import Profile

logger = get_logger()


def parse_resume_file(content: bytes, filename: str) -> str:
    """Extract raw text from PDF, DOCX, TXT, or MD bytes."""
    fname_lower = filename.lower()
    if fname_lower.endswith(".pdf"):
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(content))
        pages_text = [page.extract_text() or "" for page in reader.pages]
        return "\n\n".join(pages_text).strip()

    if fname_lower.endswith(".docx"):
        import docx

        doc = docx.Document(io.BytesIO(content))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)
        return "\n\n".join(paragraphs).strip()

    # Plain text / Markdown
    try:
        return content.decode("utf-8")
    except UnicodeDecodeError:
        return content.decode("latin-1", errors="replace")


RESUME_EXTRACTION_SYSTEM = """\
You are an expert technical resume parser and career architect.
Your task is to extract a candidate's complete career history into a structured Profile.

Guidelines:
1. Extract personal and contact details (name, email, phone, location, links).
2. Write a clear 2-3 sentence elevator pitch / sales summary in executive_summary.
3. If specific positioning guidance or directives are evident, capture in model_guidance.
4. Extract education history (institution, degree, details).
5. Categorize skills into logical groups (e.g. Infrastructure, AI Systems, Frontend, Languages).
6. In Experience, capture roles, dates, locations, optional punchy subheadings (only when
   grouping distinct project scopes; omit for single-focus roles), and impact bullets.
7. In Projects, capture standalone open-source or personal projects with optional heading and URLs.
8. Identify work authorization/citizenship and location preferences in eligibility.
9. Identify any explicit non-negotiable dealbreakers in dealbreakers.
"""


def extract_profile_from_resume_text(
    resume_text: str,
    model_name: str = DEFAULT_AGENT_MODEL,
    existing_profile: Profile | None = None,
) -> Profile:
    """Extract a Profile from raw resume text."""
    model = structured_model(model=model_name)

    user_parts = [
        "=== UPLOADED RESUME DOCUMENT ===",
        resume_text,
    ]
    if existing_profile and existing_profile.name:
        user_parts.append(
            f"\n=== EXISTING PROFILE (PRESERVE/MERGE VALUABLE DETAILS) ===\n"
            f"{existing_profile.model_dump_json(indent=2)}"
        )

    user_parts.append(
        "\nExtract and construct the comprehensive candidate Profile from this resume text."
    )

    messages = [
        ("system", RESUME_EXTRACTION_SYSTEM),
        ("user", "\n".join(user_parts)),
    ]

    extracted: Profile = invoke_structured(
        model=model,
        schema=Profile,
        messages=messages,
        label="resume_extractor",
    )
    return extracted


class CopilotChatResult(BaseModel):
    reply: str = Field(
        description="Conversational response in markdown (guidance, explanations, questions)."
    )
    updated_profile: Profile = Field(
        description="The updated Profile incorporating requested modifications."
    )
    changes_made: list[str] = Field(
        default_factory=list,
        description="Summary list of specific changes applied to the profile.",
    )


COPILOT_CHAT_SYSTEM = """\
You are CareerRadar's Profile Copilot — a technical career advisor and profile strategist.
You collaborate with the user to perfect their Profile for scoring and resume generation.

The Profile is the single source of truth for:
- 1-page tailored resume generation (executive summary, experience, projects, skills).
- Job fit scoring (experience, skills, work eligibility, model guidance, dealbreakers).

Your responsibilities:
1. Answer questions about the candidate's profile, positioning, and market readiness.
2. When asked to add, remove, or modify sections, apply changes directly to `updated_profile`.
3. Keep `executive_summary` focused as a recruiter-facing sales pitch for the resume header.
4. Keep `model_guidance` focused on internal steering directives for AI scoring and tailoring.
5. In `projects`, organize personal/open-source projects (with URLs and impact bullets).
6. In `dealbreakers`, capture hard non-negotiable rejection criteria.
7. Keep bullet points punchy, impact-focused, and quantified where possible.
8. Provide actionable, concise markdown responses.
"""


def chat_with_copilot(
    messages: list[dict[str, str]],
    current_profile: Profile,
    resume_text: str | None = None,
    model_name: str = DEFAULT_AGENT_MODEL,
) -> CopilotChatResult:
    """Run an interactive turn with the Profile Copilot."""
    model = structured_model(model=model_name)

    context_blocks = [
        "=== CURRENT CANDIDATE PROFILE ===",
        current_profile.model_dump_json(indent=2),
    ]
    if resume_text:
        context_blocks.append(f"\n=== UPLOADED RESUME TEXT CONTEXT ===\n{resume_text[:10000]}")

    formatted_messages: list[tuple[str, str]] = [
        ("system", COPILOT_CHAT_SYSTEM),
        ("user", "\n".join(context_blocks)),
    ]

    for msg in messages:
        role = msg.get("role", "user")
        content = msg.get("content", "")
        role_tag = role if role in ("system", "user", "assistant") else "user"
        formatted_messages.append((role_tag, content))

    result: CopilotChatResult = invoke_structured(
        model=model,
        schema=CopilotChatResult,
        messages=formatted_messages,
        label="profile_copilot_chat",
    )
    return result
